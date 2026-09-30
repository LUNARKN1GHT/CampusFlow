"""Generate the synthetic Word samples used by the N001 acceptance samples.

运行（在 backend/ 目录）：
    uv run --with python-docx ../samples/synthetic/n001/generate.py
或在任意已安装 python-docx 的环境中直接执行本脚本。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent


def build_paragraph_notice() -> None:
    """段落型课程通知：段落与编号列表。"""
    doc = Document()
    doc.add_heading("数据库原理 · 作业四通知", level=1)
    doc.add_paragraph("各位同学：")
    doc.add_paragraph("作业四《索引设计实验报告》现已发布，请按以下要求完成：")
    doc.add_paragraph(
        "完成 B+ 树索引的插入与删除过程推演，绘制每一步的索引结构变化图。",
        style="List Number",
    )
    doc.add_paragraph(
        "基于课程实验平台验证推演结果，并附截图与文字说明。",
        style="List Number",
    )
    doc.add_paragraph("提交方式：通过课程平台“作业四”入口提交 PDF 文件，文件名格式为“学号-姓名-作业四.pdf”。")
    doc.add_paragraph("截止时间：2026-11-15 18:00，逾期不再接收补交。")
    doc.add_paragraph("适用对象：CS2026-A 教学班，其他教学班以各自班级通知为准。")
    doc.add_paragraph("如有疑问，请通过课程讨论区提出。")
    doc.save(ROOT / "word-notice-paragraph.docx")


def build_table_notice() -> None:
    """表格型考试安排：段落说明 + 4 行 4 列表格。"""
    doc = Document()
    doc.add_heading("数据结构 · 期末考试安排（CS2026-B）", level=1)
    doc.add_paragraph("期末考试安排如下表所示，请提前 20 分钟到达考场，并携带学生证。")

    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    rows = [
        ("科目", "考试时间", "考场", "备注"),
        ("数据结构", "2026-12-20 09:00-11:00", "教三 201", "闭卷"),
        ("操作系统", "2026-12-22 14:00-16:00", "教三 302", "闭卷"),
        ("计算机网络", "2026-12-24 09:00-11:00", "教四 105", "开卷"),
    ]
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value

    doc.add_paragraph("缓考申请请于考前一周提交至学院教务办公室。")
    doc.save(ROOT / "word-notice-table.docx")


def main() -> None:
    build_paragraph_notice()
    build_table_notice()
    print("已生成 word-notice-paragraph.docx 与 word-notice-table.docx")


if __name__ == "__main__":
    main()
