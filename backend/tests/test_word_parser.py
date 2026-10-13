"""N003：真实 DOCX 样本、稳定定位及上传到 Worker 结果的闭环。"""

import io
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document

from campusflow.application import processing
from campusflow.application.materials import fragments_to_chunks
from campusflow.domain.errors import PermanentError
from campusflow.domain.word import DOCX_MIME
from campusflow.infrastructure.parsers.word import WordDocumentParser
from campusflow.infrastructure.repositories import build_repositories
from campusflow.infrastructure.storage.local import LocalFileStorage
from campusflow.workers.main import build_parser_registry

SAMPLES = Path(__file__).parents[2] / "samples/synthetic/n001"


def bytes_of(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def parse(document):
    return WordDocumentParser().parse(bytes_of(document), DOCX_MIME)


def test_sample_paragraph_numbers_match_golden_evidence():
    content = (SAMPLES / "word-notice-paragraph.docx").read_bytes()
    result = WordDocumentParser().parse(content, DOCX_MIME)
    assert not result.failures
    assert result == WordDocumentParser().parse(content, DOCX_MIME)
    assert "截止" in next(f.text for f in result.fragments if f.locator.paragraph == 7)
    assert all(f.locator.page is None for f in result.fragments)
    first = fragments_to_chunks(result, 11)
    second = fragments_to_chunks(result, 12)
    assert {item.version_id for item in first} == {11}
    assert {item.version_id for item in second} == {12}
    assert [(item.locator_type, item.locator_value) for item in first] == [
        (item.locator_type, item.locator_value) for item in second
    ]


def test_table_row_and_column_match_golden_evidence():
    result = WordDocumentParser().parse(
        (SAMPLES / "word-notice-table.docx").read_bytes(), DOCX_MIME
    )
    course = next(f for f in result.fragments if f.text == "操作系统")
    assert course.locator.as_chunk_locator() == ("table_cell", "t1:r3:c1:p1")
    row = [f for f in result.fragments if f.locator.table == 1 and f.locator.row == 3]
    assert len(row) == 4
    assert any("教三 302" in f.text for f in row)
    assert not result.failures


def test_empty_paragraphs_do_not_shift_positions_and_blocks_keep_order():
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("before")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
    doc.add_paragraph("after")
    result = parse(doc)
    assert [f.text for f in result.fragments] == ["before", "cell", "after"]
    assert result.fragments[0].locator.paragraph == 2
    assert result.fragments[-1].locator.paragraph == 3


def test_merged_and_nested_cells_have_unique_stable_locations():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    cell = table.cell(0, 0).merge(table.cell(1, 1))
    cell.text = "merged"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "nested"
    result = parse(doc)
    assert [f.text for f in result.fragments] == ["merged", "nested"]
    assert [f.locator.table for f in result.fragments] == [1, 2]


def test_image_content_is_reported_as_uncovered():
    doc = Document()
    doc.add_paragraph("visible text")
    image_path = Path(__file__).parents[2] / "samples/synthetic/d013/table-screenshot.png"
    doc.add_picture(str(image_path))
    result = parse(doc)
    assert result.is_partial
    assert "未识别" in result.failures[0].reason


def test_empty_document_is_not_false_success():
    result = parse(Document())
    assert not result.fragments and result.failures


def test_invalid_input_raises_readable_permanent_failure():
    with pytest.raises(PermanentError):
        WordDocumentParser().parse(b"fake", DOCX_MIME)


def test_uploaded_word_is_processed_and_locators_can_be_read(client, tmp_path, monkeypatch):
    monkeypatch.setattr(client.app.state, "file_storage", LocalFileStorage(tmp_path))
    monkeypatch.setattr(
        client.app.state, "job_queue", Mock(enqueue_parse=Mock(return_value="word-job"))
    )
    registry = build_parser_registry(client.app.state.settings, demo=True)
    monkeypatch.setattr(processing, "_PARSER_REGISTRY", registry)
    content = (SAMPLES / "word-notice-table.docx").read_bytes()
    response = client.post("/api/v1/materials/upload", files=[("files", ("table.docx", content))])
    assert response.status_code == 201
    item = response.json()["results"][0]
    with client.app.state.session_factory() as session:
        processing.parse_version(
            build_repositories(session), client.app.state.file_storage, item["job_id"]
        )
    mid = item["material"]["id"]
    assert client.get(f"/api/v1/materials/{mid}").json()["status"] == "done"
    chunks = client.get(f"/api/v1/materials/{mid}/versions/1/chunks").json()
    course = next(c for c in chunks if c["text"] == "操作系统")
    assert course["locator_value"] == "t1:r3:c1:p1"
    assert course["page"] is None
    assert client.get(f"/api/v1/materials/{mid}/versions/1/file").content == content
    with client.app.state.session_factory() as session:
        assert (
            processing.parse_version(
                build_repositories(session), client.app.state.file_storage, item["job_id"]
            )
            == []
        )
    assert len(client.get(f"/api/v1/materials/{mid}/versions/1/chunks").json()) == len(chunks)
